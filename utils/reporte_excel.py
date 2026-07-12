"""
Generador de reportes en formato Excel (.xlsx) usando openpyxl.
Mantiene la misma firma lógica que el generador PDF para intercambiarse
según el formato solicitado por el usuario.
"""
import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


class ExcelReporteGenerator:
    @staticmethod
    def generar_reporte_estadistico(tipo: str, periodo_inicio: str, periodo_fin: str,
                                    datos_resumen: dict, ruta_destino: str) -> str:
        """Construye una hoja de cálculo con la cabecera y los KPIs del reporte."""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = f"Reporte {tipo}"[:31]

            titulo_font = Font(size=16, bold=True, color="2A5C4D")
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="2A5C4D")
            center = Alignment(horizontal="center")

            ws["A1"] = f"REPORTE EJECUTIVO DE {tipo.upper()}"
            ws["A1"].font = titulo_font
            ws.merge_cells("A1:B1")

            ws["A2"] = f"Rango: {periodo_inicio} a {periodo_fin}"
            ws.merge_cells("A2:B2")

            ws["A4"] = "Métrica / Indicador"
            ws["B4"] = "Valor Consolidado"
            for celda in ("A4", "B4"):
                ws[celda].font = header_font
                ws[celda].fill = header_fill
                ws[celda].alignment = center

            fila = 5
            for k, v in datos_resumen.items():
                ws[f"A{fila}"] = str(k).replace("_", " ").title()
                ws[f"B{fila}"] = v
                ws[f"B{fila}"].alignment = center
                fila += 1

            ws.column_dimensions["A"].width = 40
            ws.column_dimensions["B"].width = 22

            wb.save(ruta_destino)
            return ruta_destino
        except Exception as e:
            logging.error(f"Fallo al generar reporte Excel: {str(e)}")
            raise e
