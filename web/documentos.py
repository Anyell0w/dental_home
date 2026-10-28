"""Recetas y reportes en PDF (ReportLab) y Excel (openpyxl), con la identidad de Dental Home."""
import io
import os
from datetime import datetime

from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "img", "logo.png")
MORADO, ACENTO, LAVANDA, SUAVE, LINEA = "#4A2E80", "#7B1FA2", "#B39DDB", "#F4F1F8", "#E1D5E7"
TINTA, GRIS = "#1E1433", "#5B5070"

_base = ParagraphStyle("b", fontName="Helvetica", fontSize=10, leading=14, textColor=colors.HexColor(TINTA))
_gris = ParagraphStyle("g", parent=_base, textColor=colors.HexColor(GRIS), fontSize=9, leading=12)
_titulo = ParagraphStyle("t", parent=_base, fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.HexColor(MORADO))
_marca = ParagraphStyle("m", parent=_base, fontName="Helvetica-Bold", fontSize=11, leading=13, textColor=colors.HexColor(MORADO))
_kicker = ParagraphStyle("k", parent=_base, fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.HexColor(ACENTO))
_negrita = ParagraphStyle("n", parent=_base, fontName="Helvetica-Bold")


def _esc(t) -> str:
    return (str(t) if t is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")


def _cabecera(clinica: str, titulo: str, sub: str):
    logo = Image(LOGO, width=16 * mm, height=15.5 * mm)
    marca = [Paragraph("DENTAL HOME", _marca), Paragraph(_esc(clinica), _gris)]
    t = Table([[logo, marca, [Paragraph(_esc(titulo), _titulo), Paragraph(_esc(sub), _gris)]]],
              colWidths=[20 * mm, 62 * mm, None])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (2, 0), (2, 0), "RIGHT"),
                           ("LINEBELOW", (0, 0), (-1, 0), 1.2, colors.HexColor(LINEA)), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    return t


def _tabla(filas, anchos, cabecera=True):
    t = Table(filas, colWidths=anchos, repeatRows=1 if cabecera else 0)
    estilo = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
              ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor(LINEA))]
    if cabecera:
        estilo += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(SUAVE)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(MORADO)),
                   ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
    t.setStyle(TableStyle(estilo))
    return t


def _pie(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor(GRIS))
    canvas.drawString(doc.leftMargin, 12 * mm, f"Generado con Dental Home · {datetime.now():%d/%m/%Y %H:%M}")
    canvas.drawRightString(A4[0] - doc.rightMargin, 12 * mm, f"Página {doc.page}")
    canvas.restoreState()


def _doc(buf):
    return SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=20 * mm)


def pdf_receta(clinica: str, r, meds) -> bytes:
    buf = io.BytesIO()
    doc = _doc(buf)
    numero = f"REC-{r['id']:06d}"
    story = [_cabecera(clinica, "Receta médica", f"{numero} · {r['emitida']}"), Spacer(1, 14)]
    datos = [
        [Paragraph("PACIENTE", _kicker), Paragraph("DOCUMENTO", _kicker), Paragraph("EDAD", _kicker)],
        [Paragraph(_esc(r["paciente"]), _negrita), Paragraph(f"{r['tipo_doc']} {_esc(r['numero_documento'])}", _base),
         Paragraph(f"{r['edad']} años" if r["edad"] is not None else "—", _base)],
    ]
    story += [_tabla(datos, [90 * mm, 50 * mm, None], cabecera=False), Spacer(1, 10)]
    if r["diagnostico"]:
        story += [Paragraph("DIAGNÓSTICO", _kicker), Paragraph(_esc(r["diagnostico"]), _base), Spacer(1, 12)]
    filas = [["Cant.", "Medicamento", "Indicaciones"]]
    for m in meds:
        nombre = _esc(m["nombre"]) + (f"<br/><font color='{GRIS}' size='8'>{_esc(m['presentacion'])}</font>" if m["presentacion"] else "")
        filas.append([Paragraph(f"{float(m['cantidad']):g}", _base), Paragraph(nombre, _negrita), Paragraph(_esc(m["indicaciones"]), _base)])
    story += [Paragraph("Rp/", _titulo), Spacer(1, 6), _tabla(filas, [16 * mm, 62 * mm, None])]
    if r["indicaciones_generales"]:
        story += [Spacer(1, 14), Paragraph("INDICACIONES GENERALES", _kicker), Paragraph(_esc(r["indicaciones_generales"]), _base)]
    firma = Paragraph(f"______________________________<br/><b>Dr(a). {_esc(r['doctor'])}</b><br/>"
                      f"<font color='{GRIS}'>Colegiatura {_esc(r['numero_colegiatura'] or '—')}</font>",
                      ParagraphStyle("f", parent=_base, alignment=1))
    story += [Spacer(1, 48), Table([["", firma]], colWidths=[95 * mm, None])]
    doc.build(story, onFirstPage=_pie, onLaterPages=_pie)
    return buf.getvalue()


def _grafico(series):
    d = Drawing(170 * mm, 55 * mm)
    ch = VerticalBarChart()
    ch.x, ch.y, ch.width, ch.height = 10 * mm, 8 * mm, 155 * mm, 42 * mm
    ch.data = [[v for _, v in series]]
    ch.categoryAxis.categoryNames = [k for k, _ in series]
    ch.categoryAxis.labels.fontSize = 7
    ch.valueAxis.labels.fontSize = 7
    ch.valueAxis.valueMin = 0
    ch.valueAxis.visibleGrid = True
    ch.valueAxis.gridStrokeColor = colors.HexColor(LINEA)
    ch.bars[0].fillColor = colors.HexColor(ACENTO)
    ch.bars[0].strokeColor = None
    ch.barWidth = 6
    d.add(ch)
    return d


def pdf_reporte(clinica: str, rep: dict) -> bytes:
    buf = io.BytesIO()
    doc = _doc(buf)
    story = [_cabecera(clinica, rep["titulo"], rep["periodo"]), Spacer(1, 14)]
    kpis = rep["kpis"]
    celdas = [[Paragraph(_esc(k).upper(), _kicker) for k, _ in kpis],
              [Paragraph(f"<font size='18' color='{MORADO}'><b>{_esc(v)}</b></font>", _base) for _, v in kpis]]
    story += [_tabla(celdas, None, cabecera=False), Spacer(1, 14)]
    if rep.get("serie") and any(v for _, v in rep["serie"]):
        story += [Paragraph(_esc(rep["serieTitulo"]), _negrita), Spacer(1, 4), _grafico(rep["serie"]), Spacer(1, 10)]
    for bloque in rep["tablas"]:
        story += [Paragraph(_esc(bloque["titulo"]), _negrita), Spacer(1, 6)]
        if bloque["filas"]:
            filas = [bloque["columnas"]] + [[Paragraph(_esc(c), _base) for c in f] for f in bloque["filas"]]
            story.append(_tabla(filas, None))
        else:
            story.append(Paragraph("Sin datos en el periodo.", _gris))
        story.append(Spacer(1, 14))
    doc.build(story, onFirstPage=_pie, onLaterPages=_pie)
    return buf.getvalue()


def excel_reporte(clinica: str, rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"], ws["A1"].font = rep["titulo"], Font(size=16, bold=True, color=MORADO[1:])
    ws["A2"] = f"{clinica} · {rep['periodo']}"
    fila = 4
    for k, v in rep["kpis"]:
        ws.cell(fila, 1, k).font = Font(bold=True, color=MORADO[1:])
        ws.cell(fila, 2, v)
        fila += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    cab_font, cab_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=MORADO[1:])
    for bloque in rep["tablas"]:
        hoja = wb.create_sheet(bloque["titulo"][:31])
        for j, c in enumerate(bloque["columnas"], 1):
            celda = hoja.cell(1, j, c)
            celda.font, celda.fill, celda.alignment = cab_font, cab_fill, Alignment(horizontal="center")
            hoja.column_dimensions[celda.column_letter].width = max(14, len(c) + 6)
        for i, f in enumerate(bloque["filas"], 2):
            for j, v in enumerate(f, 1):
                hoja.cell(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
